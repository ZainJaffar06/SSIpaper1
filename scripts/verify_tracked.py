# -*- coding: utf-8 -*-
"""Checks on the tracked output: (1) XML well-formed, unique revision ids, python-docx opens it;
(2) REJECT-ALL view of OUTDOC equals REJECT-ALL view of NEWDOC -> every edit is a tracked change;
(3) ACCEPT-ALL view contains the postoperative numbers and none of the old-frame numbers;
(4) citation pairing on the accepted view (check_citations)."""
import os, re, sys, zipfile, subprocess
import xml.etree.ElementTree as ET
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
NEW=os.environ['NEWDOC']; OUT=os.environ['OUTDOC']
def root(path): return ET.fromstring(zipfile.ZipFile(path).read('word/document.xml'))
def views(path):
    r=root(path); pm={c:p for p in r.iter() for c in p}
    def anc(x,tag):
        a=pm.get(x)
        while a is not None:
            if a.tag==W+tag: return True
            a=pm.get(a)
        return False
    acc=[];rej=[]
    for p in r.iter(W+'p'):
        a=''.join((t.text or '') for t in p.iter(W+'t'))          # w:t never sits inside w:del
        rj=''.join((t.text or '') for t in p.iter() if (t.tag==W+'t' or t.tag==W+'delText') and not anc(t,'ins'))
        if a.strip(): acc.append(a.strip())
        if rj.strip(): rej.append(rj.strip())
    return acc,rej,r
ok=True
r=root(OUT); ids=[e.get(W+'id') for e in r.iter() if e.tag in (W+'ins',W+'del')]
print(f'[1] well-formed; {len(ids)} revision marks; ids unique: {len(ids)==len(set(ids))}'); ok&=len(ids)==len(set(ids))
import docx; d=docx.Document(OUT); print(f'    python-docx opens: {len(d.paragraphs)} paragraphs, {len(d.tables)} tables')
accN,rejN,_=views(NEW); accO,rejO,_=views(OUT)
same=rejN==rejO; print(f'[2] reject-all(OUT) == reject-all(NEW): {same}')
if not same:
    import difflib; ok=False
    for l in list(difflib.unified_diff(rejN,rejO,lineterm='',n=0))[:30]: print('    '+l[:160])
A='\n'.join(accO)
import pandas as pd
D='/Users/zainjaffar/ssi_work/paper1_2026-08-15/out/postop_only/derived/'
MUST=['132 patients with 8 subsequent infections','0.705 (95% confidence interval [CI], 0.466–0.921)','99.6%','0.378','−3.165 to −0.259','−0.013 to +0.971',
      '0.598 (95% CI, 0.429–0.753)','0.722 (0.547–0.871)','0.702 (0.528–0.855)','+0.047±0.067','0.544 (95% CI, 0.363–0.715)','0.679 (0.507–0.829)','0.309, 0.140–0.495',
      '3,110 postoperative photographs','0.678 (95% CI, 0.536–0.805)','sensitivity was 50.0% and 22.0%','Capturing 7 of 8 infections (87.5% sensitivity) required flagging 63.6%',
      '+0.052 for EfficientNet-B0, +0.059 for ResNet18 and +0.013 for ViT-B/16','0.712 at image level and 0.729 at patient level']
OLD=['192 patients with 14','0.754 (95% confidence','98.1%','approximately 0.662','0.802 (0.680','0.727, against 0.636','13 of 14 infections','210-patient frame','206-patient','1,886','203 patients with dated','0.465 (95% CI','within 0.016','0.465 (0.327','−0.009±0.052','0.685 at image level']
miss=[m for m in MUST if m not in A]; stale=[o for o in OLD if o in A]
print(f'[3] accepted view: {len(MUST)-len(miss)}/{len(MUST)} required strings present; missing: {miss}')
print(f'    stale old-frame strings still present: {stale}'); ok&=not miss and not stale
print('[4] citations:'); out=subprocess.run([sys.executable,'-I','check_citations.py',OUT],capture_output=True,text=True).stdout
print('    '+out.split('\n')[0]); bad=[l for l in out.split('\n') if 'NO SUCH ENTRY' in l]; print(f'    citations without an entry: {len(bad)}'); ok&=not bad
print('VERIFY_TRACKED', 'PASS' if ok else 'FAIL')
